'''
Created on 2020年11月10日

@author: 20191209
'''
# -*- coding: utf-8 -*-
#from tools import get_data
from docx import Document
import xlrd
from docxtpl import DocxTemplate

data_file=r'data.xlsx'
template=r'template.docx'
result_file=r'dict.docx'

def fill_tpl(context,doc_path):
    tpl=DocxTemplate(template)
    tpl.render(context);
    tpl.save(doc_path)

def get_data(file):#
    workbook=xlrd.open_workbook(file)
#    print (workbook.sheet_names())  
    sheet2=workbook.sheet_by_name('Sheet1')  
#     nrows=sheet2.nrows  
#     ncols=sheet2.ncols  
    tables=[]
    table={}
    columns=[]
    for i in range(sheet2.nrows):
        if len(sheet2.cell(i,0).value)>0:
            if len(columns)>0:
                table['columns']=columns
                tables.append(table)
                columns=[]
            table={'table_name_cn':sheet2.cell(i,1).value,'table_name':sheet2.cell(i,0).value}
        row={'num':int(sheet2.cell(i,2).value),'name':sheet2.cell(i,3).value,'comment':sheet2.cell(i,4).value,'type':sheet2.cell(i,5).value,'typelength':int(sheet2.cell(i,6).value),'isnull':sheet2.cell(i,7).value,'ispk':sheet2.cell(i,8).value}
        columns.append(row)
       
    context={'tables':tables}
    return context  
  
if __name__ == "__main__":
    data = get_data(data_file)
    fill_tpl(data,result_file)
    print('OK.')
    